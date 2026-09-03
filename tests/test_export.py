"""Word 导出测试：Markdown → docx 的往返验证。"""
import pytest

from docx import Document

from feasibility_ai.reports import markdown_to_docx


def test_export_basic_markdown(tmp_path):
    md = (
        "# 可研报告标题\n\n"
        "## 一、项目概述\n\n"
        "这是**加粗**的一段正文。\n\n"
        "| 情景 | NPV(万元) | IRR |\n"
        "|---|---|---|\n"
        "| 悲观 | -5.2 | 6.1% |\n"
        "| 基准 | 94.5 | 46.8% |\n\n"
        "- 风险一：客流不足\n"
        "- 风险二：租金上涨\n\n"
        "> 口径声明：全投资口径。\n"
    )
    out = tmp_path / "报告.docx"
    markdown_to_docx(md, out)
    assert out.exists() and out.stat().st_size > 0

    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    assert any("可研报告标题" in t for t in texts)
    assert any("一、项目概述" in t for t in texts)
    assert any("口径声明" in t for t in texts)
    # 表格被写出
    assert len(doc.tables) == 1
    rows = [[c.text for c in row.cells] for row in doc.tables[0].rows]
    assert rows[0][0] == "情景"
    assert rows[2][1] == "94.5"


def test_export_creates_dir(tmp_path):
    out = tmp_path / "sub" / "nested" / "r.docx"
    markdown_to_docx("# x\n\n正文", out)
    assert out.exists()
