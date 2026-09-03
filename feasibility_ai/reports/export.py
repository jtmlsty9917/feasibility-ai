"""Markdown → Word(.docx) 导出（M3）。

支持报告常用的子集：标题(#~####)、表格(|)、无序列表(-)、引用(>)、加粗(**x**)。
中文字体统一宋体（SimSun），保证在 Word/WPS 打开不乱码、格式统一。
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm

CN_FONT = "宋体"


def _set_cn_font(doc: Document) -> None:
    """把默认字体与东亚字体都设为宋体。"""
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2",
                       "Heading 3", "Heading 4", "List Bullet"):
        try:
            style = doc.styles[style_name]
            style.font.name = CN_FONT
            style.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
        except Exception:
            pass


def _add_runs(paragraph, text: str) -> None:
    """支持 **加粗** 的行内标记。"""
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        if i % 2 == 1:
            run.bold = True


def _is_separator_row(line: str) -> bool:
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c) for c in cells)


def markdown_to_docx(md_text: str, out_path: str | Path) -> Path:
    """把 Markdown 文本写成一个 Word 文档，返回输出路径。"""
    doc = Document()
    _set_cn_font(doc)

    lines = md_text.splitlines()
    table_buffer: list[list[str]] = []
    in_code = False

    def flush_table() -> None:
        if not table_buffer:
            return
        rows = table_buffer
        n_cols = max(len(r) for r in rows)
        table = doc.add_table(rows=len(rows), cols=n_cols)
        table.style = "Light Grid Accent 1"
        for ri, row in enumerate(rows):
            for ci in range(n_cols):
                cell_text = row[ci] if ci < len(row) else ""
                _add_runs(table.cell(ri, ci).paragraphs[0], cell_text)
        doc.add_paragraph()  # 表后空行
        table_buffer.clear()

    first_content = True
    for raw in lines:
        line = raw.rstrip()

        # 代码块跳过（报告一般无代码，防御用）
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            continue

        # 表格行
        if line.strip().startswith("|"):
            if _is_separator_row(line):
                continue  # 分隔行不建行
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            table_buffer.append(cells)
            continue
        flush_table()

        stripped = line.strip()
        if not stripped:
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = min(len(m.group(1)), 4)
            heading = doc.add_heading(level=level)
            _add_runs(heading, m.group(2).strip())
            first_content = False
            continue
        if stripped.startswith((">", "＞")):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)  # 引用缩进
            run = p.add_run(stripped.lstrip(">").strip())
            run.italic = True
            continue
        if re.match(r"^[-*]\s+", stripped):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, re.sub(r"^[-*]\s+", "", stripped))
            continue

        p = doc.add_paragraph()
        _add_runs(p, stripped)

    flush_table()

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out
